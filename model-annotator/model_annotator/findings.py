"""Findings engine (pipeline Phases 5-7).

Phase 5 — candidate rules fire over metrics + structure + the formula graph.
Phase 6 — every candidate gets an acquittal attempt (the known decompositions)
          before it may become a finding; acquitted candidates are reported
          explicitly, with their resolution.
Phase 7 — severity, confidence caps, cause-vs-echo dedup (echoes collapse onto
          the finding closest to the source), and trust degradation.

Tone contract: findings surface tensions and ask neutral management questions.
They never accuse — every pattern here can have a legitimate explanation.
"""
from __future__ import annotations

import logging
import re
from dataclasses import dataclass, field
from typing import Optional

from .graph import FormulaGraph, errors_feeding_sheets, orphan_formula_rows
from .ingest import WorkbookData, is_number
from .integrity import IntegrityResult
from .mapping import LineItem, MappingResult
from .metrics import EPS, MetricsResult, find_scalar_assumptions
from .schema import (
    Acquittal,
    AssumptionEntry,
    CellCitation,
    CleanCheck,
    Finding,
    QuantifiedImpact,
    Severity,
    SheetRole,
    TieOutStatus,
    format_ref,
    parse_ref,
)
from .structure import StructureResult

log = logging.getLogger(__name__)

SMALL_BASE_FRACTION = 0.01
GRANT_DEPENDENCE_THRESHOLD = 0.30
CONCENTRATION_THRESHOLD = 0.50
CLIFF_MONTHS = 3.0


@dataclass
class Candidate:
    category: str
    severity: Severity
    title: str
    narrative: str
    management_question: str
    evidence: list[CellCitation] = field(default_factory=list)
    confidence: float = 0.8
    quantified_impact: Optional[QuantifiedImpact] = None
    periods: list[str] = field(default_factory=list)
    acquit_with: tuple[str, ...] = ()        # decomposition ids to attempt
    mapping_confidence: float = 1.0          # min confidence of mappings relied on

    def cells(self) -> list[str]:
        return [e.ref for e in self.evidence]


@dataclass
class FindingsResult:
    findings: list[Finding] = field(default_factory=list)
    acquittals: list[Acquittal] = field(default_factory=list)
    clean_checks: list[CleanCheck] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class FCtx:
    wbd: WorkbookData
    structure: StructureResult
    mapping: MappingResult
    graph: Optional[FormulaGraph]
    integrity: IntegrityResult
    metrics: MetricsResult
    census: list[AssumptionEntry]
    benchmarks: dict

    @property
    def sheet(self) -> Optional[str]:
        return self.mapping.primary_sheet

    def periods(self) -> list[str]:
        ax = self.structure.primary_axis(self.sheet) if self.sheet else None
        return list(ax.periods) if ax else []

    def item(self, cid: str) -> Optional[LineItem]:
        return self.mapping.get(cid, self.sheet)

    def cite_series(self, it: Optional[LineItem], periods: list[str]) -> list[CellCitation]:
        out = []
        for p in periods:
            if it is None:
                break
            ref = it.ref_for(p)
            if ref:
                out.append(CellCitation(ref=ref, value=it.value_for(p), label=f"{it.label.strip()} [{p}]"))
        return out


# ===========================================================================
# Phase 5 — candidate rules
# ===========================================================================

def _rule_tie_out_failures(ctx: FCtx, cands: list[Candidate]) -> None:
    for t in ctx.integrity.tie_outs:
        if t.status == TieOutStatus.failed and t.material:
            cands.append(Candidate(
                category="arithmetic_integrity",
                severity=Severity.critical,
                title=f"The model violates its own arithmetic: {t.label}",
                narrative=(f"Tie-out '{t.id}' fails: {t.detail} A model that breaks its own identities "
                           "cannot be trusted on anything downstream of the broken rows."),
                management_question=("Can you walk us through the cash build in the failing periods? "
                                     "We want to make sure we are reading the statement structure as intended."),
                evidence=t.evidence,
                confidence=0.95,
            ))
        elif t.status == TieOutStatus.unfalsifiable:
            cands.append(Candidate(
                category="unfalsifiable_check",
                severity=Severity.critical,
                title="The model's cash check cannot fail by construction",
                narrative=(f"{t.detail} Never running out of money is a mechanism here, not a result: the "
                           "financing line is computed to absorb any shortfall, so the check carries no "
                           "information about financing risk."),
                management_question=("If the auto-raise mechanism is switched off, in which months does cash "
                                     "go negative, and how large is the gap? What is the concrete plan for "
                                     "the raises the model inserts automatically?"),
                evidence=t.evidence,
                confidence=0.9,
            ))


def _rule_missing_cached(ctx: FCtx, cands: list[Candidate]) -> None:
    if ctx.wbd.missing_cached_values:
        cands.append(Candidate(
            category="data_quality",
            severity=Severity.critical,
            title="Workbook saved without calculated values",
            narrative=(f"{ctx.wbd.n_formula_cells_without_cached_value:,} of {ctx.wbd.n_formula_cells:,} "
                       "formula cells carry no cached value, so their numbers cannot be read without "
                       "recalculating in Excel. All numeric analysis here is partial."),
            management_question="Could you re-send the model saved after a full calculation pass?",
            confidence=0.95,
        ))


def _rule_units_inconsistency(ctx: FCtx, cands: list[Candidate]) -> None:
    for t in ctx.integrity.tie_outs:
        if t.id.startswith("mirror_") and t.status == TieOutStatus.failed and "units inconsistency" in t.detail:
            cands.append(Candidate(
                category="unit_scale_inconsistency",
                severity=Severity.high,
                title="Sheets disagree on units by a consistent factor",
                narrative=t.detail,
                management_question=("Which sheets are in thousands and which in whole dollars? "
                                     "A mixed-units link can silently inflate or deflate a statement 1000x."),
                evidence=t.evidence,
                confidence=0.85,
            ))


def _rule_grant_dependence(ctx: FCtx, cands: list[Candidate]) -> None:
    share = ctx.metrics.series("grant_share_of_revenue")
    if not share:
        return
    P = ctx.periods()
    heavy = [p for p in P if (share.get(p) or 0) > GRANT_DEPENDENCE_THRESHOLD]
    if not heavy:
        return
    worst_p = max(heavy, key=lambda p: share[p])
    m = ctx.metrics.get("grant_share_of_revenue")
    ex_growth = ctx.metrics.series("revenue_ex_grants_growth")
    rev_growth = ctx.metrics.series("revenue_growth")
    hidden_decline = [p for p in P
                      if (ex_growth.get(p) is not None and ex_growth[p] < -0.02)
                      and (share.get(p) or 0) > 0.2]
    narrative = (f"Grants supply more than {GRANT_DEPENDENCE_THRESHOLD:.0%} of revenue in "
                 f"{len(heavy)} period(s) (peak {share[worst_p]:.0%} in {worst_p}). Grant revenue typically "
                 "carries 100% margin and is not commercially repeatable.")
    sev = Severity.high
    if hidden_decline:
        dp = hidden_decline[0]
        tot = rev_growth.get(dp)
        narrative += (f" More importantly, commercial (ex-grant) revenue actually DECLINES "
                      f"{ex_growth[dp]:.0%} in {dp}"
                      + (f" while total revenue growth reads {tot:+.0%}" if tot is not None else "")
                      + " — the grant line smooths over a commercial setback.")
    rev_item = ctx.item("revenue_total")
    grant_items = [it for it in ctx.mapping.get_all("grant_or_subsidy_revenue") if it.sheet == ctx.sheet]
    ev = []
    for it in grant_items[:2]:
        ev += ctx.cite_series(it, heavy[:2] + hidden_decline[:1])
    ev += ctx.cite_series(rev_item, [worst_p] + hidden_decline[:1])
    cands.append(Candidate(
        category="grant_dependence",
        severity=sev,
        title=("Grant revenue smooths a commercial decline" if hidden_decline
               else f"Grants carry up to {share[worst_p]:.0%} of revenue"),
        narrative=narrative,
        management_question=("How much of the granted amount is awarded versus applied-for, and what are the "
                             "milestones attached? How would the plan change if the next tranche slipped a year?"),
        evidence=ev,
        confidence=0.85,
        quantified_impact=QuantifiedImpact(
            metric="revenue_ex_grants_growth",
            as_modeled=rev_growth.get(hidden_decline[0]) if hidden_decline else None,
            as_corrected=ex_growth.get(hidden_decline[0]) if hidden_decline else None,
            basis="total revenue growth vs growth excluding grant rows, model numbers only",
        ) if hidden_decline else None,
        periods=heavy,
        mapping_confidence=min((it.confidence for it in grant_items), default=1.0),
    ))


def _rule_growth_whiplash(ctx: FCtx, cands: list[Candidate]) -> None:
    P = ctx.periods()
    for mid, label in (("revenue_growth", "Revenue"), ("revenue_ex_grants_growth", "Revenue ex-grants")):
        gs = ctx.metrics.series(mid)
        if not gs:
            continue
        vals = [(p, gs[p]) for p in P if gs.get(p) is not None]
        for (p1, g1), (p2, g2) in zip(vals, vals[1:]):
            swing = (g1 > 0.5 and g2 < -0.2) or (g1 < -0.2 and g2 > 0.5) or \
                    (abs(g1) > EPS and abs(g2 / g1) > 5 and abs(g2) > 1.0)
            if not swing:
                continue
            it = ctx.item("revenue_total")
            cands.append(Candidate(
                category="growth_whiplash",
                severity=Severity.medium,
                title=f"{label} growth whipsaws between {p1} ({g1:+.0%}) and {p2} ({g2:+.0%})",
                narrative=(f"{label} growth swings from {g1:+.0%} to {g2:+.0%} across consecutive periods. "
                           "Swings like this usually trace to one-off events (a grant, a single project) "
                           "rather than the underlying business."),
                management_question=f"What specifically drives the {p2} swing — new contracts, timing, or a one-off?",
                evidence=ctx.cite_series(it, [p1, p2]),
                confidence=0.7,
                periods=[p1, p2],
                acquit_with=("small_base", "grant_timing"),
            ))
            break  # one whiplash candidate per series tells the story


def _rule_hockey_stick(ctx: FCtx, cands: list[Candidate]) -> None:
    P = ctx.periods()
    gs = ctx.metrics.series("revenue_growth")
    rev = ctx.metrics.get("revenue_growth")
    if not gs or rev is None:
        return
    it = ctx.item("revenue_total")
    if it is None:
        return
    levels = {p: it.value_for(p) for p in P}
    scale = max((abs(v) for v in levels.values() if v is not None), default=0.0)
    runs: list[list[str]] = []
    cur: list[str] = []
    for p in P:
        g = gs.get(p)
        lvl = levels.get(p)
        if g is not None and g > 1.0 and lvl is not None and abs(lvl) > 0.05 * scale:
            cur.append(p)
        else:
            if len(cur) >= 2:
                runs.append(cur)
            cur = []
    if len(cur) >= 2:
        runs.append(cur)
    for run in runs[:1]:
        gtxt = ", ".join(f"{p}: {gs[p]:+.0%}" for p in run)
        cands.append(Candidate(
            category="hockey_stick",
            severity=Severity.medium,
            title=f"Sustained >100% growth into large bases ({run[0]}–{run[-1]})",
            narrative=(f"Revenue more than doubles in consecutive periods that are already at scale ({gtxt}). "
                       "Compounding at this rate from a large base is the strongest claim in the model."),
            management_question=("What has to be true operationally — capacity, sales cycle, hiring — for "
                                 f"the {run[0]}–{run[-1]} doubling years to land?"),
            evidence=ctx.cite_series(it, run),
            confidence=0.75,
            periods=run,
        ))


def _rule_financing_cliff(ctx: FCtx, cands: list[Candidate]) -> None:
    P = ctx.periods()
    fwd = ctx.metrics.series("runway_forward_months")
    if not fwd:
        return
    cliffs = [p for p in P if fwd.get(p) is not None and 0 <= fwd[p] < CLIFF_MONTHS]
    if not cliffs:
        return
    p0 = cliffs[0]
    same = ctx.metrics.series("runway_same_period_months")
    eq = ctx.item("equity_proceeds")
    rescue = None
    if eq is not None:
        idx = P.index(p0)
        for p in P[idx + 1: idx + 3]:
            v = eq.value_for(p)
            if v is not None and v > EPS:
                ref = eq.ref_for(p)
                sheet_name, cell = parse_ref(ref)
                rec_sheet = ctx.wbd.sheets.get(sheet_name)
                typed = False
                if rec_sheet is not None:
                    from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
                    cl, row = coordinate_from_string(cell)
                    rec = rec_sheet.cell(row, column_index_from_string(cl))
                    typed = rec is not None and rec.formula is None
                rescue = (p, v, ref, typed)
                break
    end_it = ctx.item("ending_cash")
    narrative = (f"Forward-looking runway falls to {fwd[p0]:.1f} months entering {p0}"
                 + (f" (the same-year figure reads {same[p0]:.1f} months — the forward number is the honest "
                    f"one when budgets step up)" if same.get(p0) is not None and same[p0] > fwd[p0] * 1.5 else "")
                 + ".")
    if rescue:
        p, v, ref, typed = rescue
        narrative += (f" The model survives because a {'typed-in' if typed else 'computed'} raise of "
                      f"{v:,.0f} lands in {p}.")
    cands.append(Candidate(
        category="financing_cliff",
        severity=Severity.high,
        title=f"Runway compresses to {fwd[p0]:.1f} months before the next raise",
        narrative=narrative,
        management_question=("What is the financing plan if the round modeled "
                             + (f"in {rescue[0]} " if rescue else "") +
                             "slips two quarters or prices down — which spend gets cut first?"),
        evidence=(ctx.cite_series(end_it, [p0]) +
                  ([CellCitation(ref=rescue[2], value=rescue[1], label="rescuing raise"
                                 + (" (typed input)" if rescue[3] else ""))] if rescue else [])),
        confidence=0.85,
        periods=cliffs,
    ))


def _rule_concentration(ctx: FCtx, cands: list[Candidate]) -> None:
    term = ctx.metrics.get("terminal_concentration")
    if term is None or term.scalar is None or term.scalar < CONCENTRATION_THRESHOLD:
        return
    seg_name = (term.notes or "").replace("segment: ", "")
    P = ctx.periods()
    seg_share = ctx.metrics.series(f"segment_share::{seg_name}")
    first_active = next((p for p in P if (seg_share.get(p) or 0) > 0.01), None)
    is_new = first_active is not None and P.index(first_active) > len(P) * 2 // 3
    cands.append(Candidate(
        category="concentration",
        severity=Severity.high if is_new else Severity.medium,
        title=f"{term.scalar:.0%} of terminal revenue rides on {seg_name!r}",
        narrative=(f"The largest segment carries {term.scalar:.0%} of final-period revenue."
                   + (f" It only becomes active in {first_active}, so the model's end-state rests on its "
                      "newest, least proven line." if is_new else "")),
        management_question=(f"What is the current commercial evidence behind {seg_name!r} — signed contracts, "
                             "LOIs, or pipeline?"),
        evidence=[CellCitation(ref=r, value=None) for r in term.inputs[:6]],
        confidence=0.75,
    ))


def _rule_market_share_basis(ctx: FCtx, cands: list[Candidate]) -> None:
    printed_it = ctx.item("market_share")
    if printed_it is None:
        return
    P = ctx.periods()
    printed = {p: printed_it.value_for(p) for p in P}
    bases: dict[str, dict] = {}
    for m in ctx.metrics.metrics:
        if m.metric_id.startswith("market_share_recomputed::"):
            bases[m.metric_id.split("::", 1)[1]] = m.series
    if len(bases) < 2:
        return
    # which basis does the printed row equal, and how far apart are the bases?
    last = next((p for p in reversed(P) if printed.get(p)), None)
    if last is None:
        return
    pv = printed[last]
    matched = min(bases.items(), key=lambda kv: abs((kv[1].get(last) or 1e9) - pv))
    others = {k: v.get(last) for k, v in bases.items() if k != matched[0] and v.get(last) is not None}
    if not others or abs((matched[1].get(last) or 0) - pv) > abs(pv) * 0.2:
        return
    alt_name, alt_v = max(others.items(), key=lambda kv: abs(kv[1] - pv))
    if alt_v is None or abs(alt_v - pv) < abs(pv) * 0.3:
        return
    cands.append(Candidate(
        category="market_share_basis_mismatch",
        severity=Severity.high,
        title=f"Market share is quoted on the {matched[0]} basis ({pv:.1%}); on {alt_name} it reads {alt_v:.1%}",
        narrative=(f"The printed market-share row matches the {matched[0]} capacity basis ({pv:.1%} in {last}), "
                   f"but the model also carries a {alt_name} basis on which the same share is {alt_v:.1%}. "
                   "Quoting the smaller number next to the larger capacity headline understates implied "
                   "market-capture risk."),
        management_question=(f"Which basis — {matched[0]} or {alt_name} — should we underwrite market share on, "
                             "and why is the printed figure on the other one?"),
        evidence=[CellCitation(ref=printed_it.ref_for(last), value=pv, label="printed market share")],
        confidence=0.8,
        quantified_impact=QuantifiedImpact(
            metric="market_share", as_modeled=pv, as_corrected=alt_v,
            basis=f"recomputed on the {alt_name} capacity basis from the model's own rows"),
    ))


def _rule_frozen_costs(ctx: FCtx, cands: list[Candidate]) -> None:
    P = ctx.periods()
    rev_it = ctx.item("revenue_total")
    if rev_it is None or len(P) < 5:
        return
    rev = {p: rev_it.value_for(p) for p in P}
    valid = [p for p in P if rev.get(p) is not None and rev[p] > EPS]
    if len(valid) < 5 or rev[valid[-1]] / max(rev[valid[0]], EPS) < 10:
        return
    span = f"{valid[0]}–{valid[-1]}"
    growth_x = rev[valid[-1]] / max(rev[valid[0]], EPS)
    comps = [it for it in ctx.mapping.get_all("opex_component") if it.sheet == ctx.sheet]
    frozen: list[LineItem] = []
    for it in comps:
        vals = [it.value_for(p) for p in valid]
        vals = [v for v in vals if v is not None]
        if len(vals) >= 5 and max(vals) > EPS and (max(vals) - min(vals)) / max(abs(max(vals)), EPS) < 0.02:
            frozen.append(it)
    if not frozen:
        return
    names = ", ".join(it.label.strip() for it in frozen[:4])
    ev = []
    for it in frozen[:3]:
        ev += ctx.cite_series(it, [valid[0], valid[-1]])
    ev += ctx.cite_series(rev_it, [valid[0], valid[-1]])
    cands.append(Candidate(
        category="frozen_support_costs",
        severity=Severity.high,
        title=f"Support costs frozen while revenue grows {growth_x:,.0f}x",
        narrative=(f"{names} are held flat across {span} while revenue grows {growth_x:,.0f}x. "
                   "Someone has to answer the phones: flat support/G&A against exponential revenue "
                   "is a margin subsidy hiding in opex."),
        management_question=f"What scales {names} — and why does the model hold them constant through {span}?",
        evidence=ev,
        confidence=0.8,
        mapping_confidence=min((it.confidence for it in frozen), default=1.0),
    ))


def _rule_expense_decoupling(ctx: FCtx, cands: list[Candidate]) -> None:
    P = ctx.periods()
    rg = ctx.metrics.series("revenue_growth")
    og = ctx.metrics.series("opex_growth")
    if not rg or not og:
        return
    for p in P:
        r, o = rg.get(p), og.get(p)
        if r is not None and o is not None and r > 1.0 and o < -0.05:
            opex_it = ctx.item("opex_total")
            cands.append(Candidate(
                category="expense_growth_decoupled",
                severity=Severity.high,
                title=f"Total expenses FALL {o:.0%} in a {r:+.0%} revenue year ({p})",
                narrative=(f"In {p} revenue grows {r:+.0%} while total expenses decline {o:.0%}. "
                           "Expenses falling in a breakout year is either a structural story (one-time "
                           "spend rolling off) or a modeling error."),
                management_question=f"What rolls off in {p} that lets expenses fall while revenue multiplies?",
                evidence=ctx.cite_series(opex_it, [p]) + ctx.cite_series(ctx.item("revenue_total"), [p]),
                confidence=0.8,
                periods=[p],
                acquit_with=("ex_capex",),
            ))
            break


def _rule_capex_in_opex(ctx: FCtx, cands: list[Candidate]) -> None:
    capex_comps = [it for it in ctx.mapping.get_all("opex_component")
                   if "capex_expensed" in it.subtags and it.sheet == ctx.sheet]
    if not capex_comps:
        return
    dep = ctx.mapping.get("depreciation")
    orphan_dep = False
    if ctx.graph is not None and dep is not None:
        orphans = {(s, r) for s, r, _ in orphan_formula_rows(ctx.graph)}
        orphan_dep = (dep.sheet, dep.index) in orphans
    ebitda_before = ctx.metrics.get("ebitda_before_expensed_capex")
    names = ", ".join(it.label.strip() for it in capex_comps[:3])
    P = ctx.periods()
    last = P[-1] if P else None
    qi = None
    if ebitda_before is not None and last:
        eb = ctx.metrics.series("ebitda_margin")
        # restated margin from the restated EBITDA
        rev_it = ctx.item("revenue_total")
        if rev_it is not None and ebitda_before.series.get(last) is not None:
            rv = rev_it.value_for(last)
            if rv:
                qi = QuantifiedImpact(
                    metric=f"ebitda_margin[{last}]",
                    as_modeled=eb.get(last),
                    as_corrected=ebitda_before.series[last] / rv,
                    basis="EBITDA with expensed capital equipment added back, from the model's own rows")
    narrative = (f"Capital-equipment spend ({names}) is expensed inside operating expenses rather than "
                 "capitalized.")
    if dep is None:
        narrative += " No depreciation line exists, consistent with nothing being capitalized."
    elif orphan_dep:
        narrative += (" A depreciation schedule exists but feeds nothing downstream — computed and "
                      "orphaned while capex runs through opex.")
    narrative += (" This depresses reported EBITDA in build years and makes expense growth lumpy "
                  "(deceleration that is really capex timing).")
    ev = []
    for it in capex_comps[:2]:
        nz = [p for p in P if (it.value_for(p) or 0) > EPS][:2]
        ev += ctx.cite_series(it, nz)
    if dep is not None and orphan_dep:
        ev.append(CellCitation(ref=format_ref(dep.sheet, f"row {dep.index}"),
                               value=None, label=f"orphaned: {dep.label.strip()}"))
    cands.append(Candidate(
        category="capex_expensed_in_opex" + ("_orphaned_depreciation" if orphan_dep else ""),
        severity=Severity.high if orphan_dep else Severity.medium,
        title=("Capex expensed in opex while a depreciation schedule sits orphaned" if orphan_dep
               else "Capital equipment is expensed inside opex"),
        narrative=narrative,
        management_question=("Is the intent cash accounting? If so, which EBITDA should we underwrite — "
                             "as printed, or before expensed equipment?"),
        evidence=ev,
        confidence=0.85,
        quantified_impact=qi,
    ))


def _rule_opex_pegged_to_revenue(ctx: FCtx, cands: list[Candidate]) -> None:
    if ctx.graph is None:
        return
    opex = ctx.item("opex_total")
    rev = ctx.item("revenue_total")
    if opex is None or rev is None or opex.n_formula == 0:
        return
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    from .graph import decode
    # match by (sheet, row), not exact cell: in yearly-left/monthly-right
    # layouts the chain runs through the revenue row's MONTHLY cells, which
    # are siblings (not ancestors) of the annual cells we mapped
    rev_rows: set[tuple[int, int]] = set()
    for ref in rev.refs:
        s, cell = parse_ref(ref)
        cl, row = coordinate_from_string(cell)
        rev_rows.add((ctx.graph.sheet_index[s], row))
    pegged = 0
    checked = 0
    for ref in opex.refs[-6:]:
        s, cell = parse_ref(ref)
        cl, row = coordinate_from_string(cell)
        node = ctx.graph.node_of(s, row, column_index_from_string(cl))
        ups = ctx.graph.upstream(node, cap=120_000)
        checked += 1
        if any(n >= 0 and (decode(n)[0], decode(n)[1]) in rev_rows for n in ups):
            pegged += 1
    if checked and pegged / checked > 0.5:
        pct = find_scalar_assumptions(ctx.wbd, ctx.structure,
                                      r"%\s*of\s*revenue|percent\s*of\s*revenue", (0.0, 1.0))
        ev = [CellCitation(ref=r, value=v, label=lab) for r, v, lab in pct[:3]]
        cands.append(Candidate(
            category="opex_pegged_to_revenue",
            severity=Severity.medium,
            title="Late-year opex is defined as a percentage of revenue",
            narrative=("Operating expenses are computed from revenue itself"
                       + (f" (typed set-standards: {', '.join(f'{lab} = {v:g}' for _, v, lab in pct[:3])})"
                          if pct else "")
                       + ". Spending defined as a slice of revenue can never look bad against revenue — "
                       "the opex/revenue ratio is an input, not a result."),
            management_question=("What does the actual organization plan cost in the out-years, built "
                                 "bottom-up rather than as a percent-of-revenue standard?"),
            evidence=ev + [CellCitation(ref=opex.refs[-1], value=opex.values[-1] if opex.values else None,
                                        label=opex.label.strip())],
            confidence=0.75,
        ))


def _rule_flat_inputs(ctx: FCtx, cands: list[Candidate]) -> None:
    flat = [e for e in ctx.census
            if e.flat_all_periods
            and re.search(r"price|asp|\$/|cost per|rate|rent|salary|commodity|fee", e.label, re.I)
            and not re.search(r"tax|discount", e.label, re.I)]
    if not flat:
        return
    top = flat[:5]
    names = "; ".join(f"{e.label} = {e.value} ({e.cell})" for e in top)
    cands.append(Candidate(
        category="flat_pricing_inputs",
        severity=Severity.medium,
        title=f"{len(flat)} price/cost input(s) held flat for the entire horizon",
        narrative=(f"High-impact typed inputs are constant across every period: {names}. "
                   "Decade-flat prices and rates are a modeling convenience, and usually a flattering one "
                   "when costs decline elsewhere in the model."),
        management_question="Which of these flat inputs have contracted support, and which are placeholders?",
        evidence=[CellCitation(ref=e.cell, value=e.value, label=e.label) for e in top],
        confidence=0.7,
    ))


def _rule_one_way_margin(ctx: FCtx, cands: list[Candidate]) -> None:
    """Cost-side factors decline while price-side factors stay pinned."""
    decl_cost: list[tuple[LineItem, float]] = []
    flat_price: list[LineItem] = []
    for it in ctx.mapping.items:
        role, _ = ctx.structure.roles.get(it.sheet, (SheetRole.unknown, 0.0))
        if role not in (SheetRole.assumptions, SheetRole.cost_buildup, SheetRole.schedule):
            continue
        if it.n_formula > it.n_numeric / 2:
            continue  # typed assumption rows only
        vals = [v for v in it.values if v is not None]
        if len(vals) < 4 or max(abs(v) for v in vals) < EPS:
            continue
        lab = it.label.lower()
        is_cost = bool(re.search(r"cost|learning|decline|reduction|efficiency", lab))
        is_price = bool(re.search(r"price|asp", lab)) and not is_cost
        rel_change = (vals[-1] - vals[0]) / max(abs(vals[0]), EPS)
        spread = (max(vals) - min(vals)) / max(abs(max(vals)), EPS)
        if is_cost and rel_change < -0.10:
            decl_cost.append((it, rel_change))
        elif is_price and spread < 0.02:
            flat_price.append(it)
    if not decl_cost or not flat_price:
        return
    c_it, c_chg = decl_cost[0]
    p_it = flat_price[0]
    cands.append(Candidate(
        category="one_way_margin_expansion",
        severity=Severity.medium,
        title="Costs are assumed to fall while prices are assumed not to",
        narrative=(f"Cost-side assumption {c_it.label.strip()!r} declines {c_chg:.0%} over the horizon while "
                   f"price-side assumption {p_it.label.strip()!r} is held flat. Margin expansion built this "
                   "way is one-directional: every learning effect accrues to the company, none to customers."),
        management_question=("What competitive dynamic lets cost declines stay out of pricing for the "
                             "entire model horizon?"),
        evidence=[CellCitation(ref=c_it.refs[0], value=c_it.values[0], label=c_it.label.strip()),
                  CellCitation(ref=c_it.refs[-1], value=c_it.values[-1], label=c_it.label.strip()),
                  CellCitation(ref=p_it.refs[0], value=p_it.values[0], label=p_it.label.strip())],
        confidence=0.7,
        mapping_confidence=min(c_it.confidence or 1.0, p_it.confidence or 1.0, 1.0),
    ))


def _rule_taxes(ctx: FCtx, cands: list[Candidate]) -> None:
    refunds = ctx.metrics.get("tax_refunds_in_loss_years")
    if refunds is not None:
        ps = [p for p, v in refunds.series.items() if v is not None]
        if ps:
            tx = ctx.item("taxes")
            cands.append(Candidate(
                category="tax_logic",
                severity=Severity.medium,
                title=f"Taxes booked as income in loss years ({', '.join(ps[:4])})",
                narrative=("The model credits tax benefits in pre-tax loss years. Early-stage companies "
                           "rarely monetize losses in cash; without a refund mechanism these credits "
                           "overstate cash."),
                management_question="What refund mechanism (carryback, refundable credits) supports cash tax benefits in loss years?",
                evidence=ctx.cite_series(tx, ps[:3]),
                confidence=0.75,
                acquit_with=("tax_trend",),
            ))
    nol = ctx.metrics.get("cash_taxes_despite_cumulative_losses")
    if nol is not None:
        ps = [p for p, v in nol.series.items() if v is not None]
        if ps:
            tx = ctx.item("taxes")
            cands.append(Candidate(
                category="tax_logic",
                severity=Severity.medium,
                title=f"Full cash taxes while cumulative pre-tax income is still negative ({ps[0]}–{ps[-1]})",
                narrative=("The model pays cash taxes in periods where cumulative pre-tax income is still "
                           "negative — NOL carryforwards would normally shelter these years. This is "
                           "conservative for cash, but it means the tax line is mechanical rather than "
                           "planned."),
                management_question="Does the company expect to use NOL carryforwards, and are they modeled anywhere?",
                evidence=ctx.cite_series(tx, ps[:3]),
                confidence=0.7,
                acquit_with=("tax_trend",),
            ))


def _rule_orphans_and_errors(ctx: FCtx, cands: list[Candidate]) -> None:
    if ctx.graph is None:
        return
    stmt_sheets = [ctx.sheet] if ctx.sheet else []
    errs = errors_feeding_sheets(ctx.graph, stmt_sheets)
    feeding = [ref for ref, feeds in errs if feeds]
    dead = [ref for ref, feeds in errs if not feeds]
    if feeding:
        cands.append(Candidate(
            category="error_cells_feed_statements",
            severity=Severity.high,
            title=f"{len(feeding)} error cell(s) feed the financial statements",
            narrative=(f"Cells carrying Excel errors ({', '.join(feeding[:5])}"
                       + (", …" if len(feeding) > 5 else "") +
                       ") sit upstream of the statements. Anything downstream of an error cell is "
                       "whatever Excel last managed to compute."),
            management_question="Are these error cells known issues, and what should those branches compute?",
            evidence=[CellCitation(ref=r, value=None) for r in feeding[:6]],
            confidence=0.9,
        ))
    elif dead:
        cands.append(Candidate(
            category="error_cells_dead_branch",
            severity=Severity.info,
            title=f"{len(dead)} error cell(s) in dead branches",
            narrative=f"Errors exist ({', '.join(dead[:5])}) but nothing downstream consumes them.",
            management_question="",
            evidence=[CellCitation(ref=r, value=None) for r in dead[:6]],
            confidence=0.8,
        ))
    # orphaned computed rows on calc sheets (presentation/scratch excluded)
    orphans = orphan_formula_rows(ctx.graph, min_cells=4)
    interesting = []
    for sheet, row, n in orphans:
        role, _ = ctx.structure.roles.get(sheet, (SheetRole.unknown, 0.0))
        if role in (SheetRole.presentation, SheetRole.documentation, SheetRole.scratch):
            continue
        sd = ctx.wbd.sheets[sheet]
        lab = ""
        for cc in range(1, 6):
            rec = sd.cell(row, cc)
            if rec is not None and isinstance(rec.value, str):
                lab = rec.value.strip()
                break
        if lab and not re.search(r"check|chart|note", lab, re.I):
            interesting.append((sheet, row, n, lab))
    if len(interesting) >= 3:
        sample = "; ".join(f"{format_ref(s, f'row {r}')} ({lab[:30]!r}, {n} cells)"
                           for s, r, n, lab in interesting[:4])
        cands.append(Candidate(
            category="orphaned_computations",
            severity=Severity.info,
            title=f"{len(interesting)} computed row(s) feed nothing",
            narrative=(f"Blocks of formulas are calculated but never consumed: {sample}. Orphans are "
                       "either leftovers or — when they duplicate live logic with different answers — "
                       "a sign the wired-in version was chosen over the careful one."),
            management_question="",
            evidence=[CellCitation(ref=format_ref(s, f"row {r}"), value=None, label=lab)
                      for s, r, n, lab in interesting[:5]],
            confidence=0.7,
        ))


def _rule_circulars(ctx: FCtx, cands: list[Candidate]) -> None:
    if ctx.graph is None or not ctx.graph.circular_groups:
        return
    solver_found = any(t.status == TieOutStatus.unfalsifiable for t in ctx.integrity.tie_outs)
    if solver_found:
        return  # the circularity is evidence under the unfalsifiable finding
    grp = ctx.graph.circular_groups[0]
    cands.append(Candidate(
        category="circular_references",
        severity=Severity.medium,
        title=f"{len(ctx.graph.circular_groups)} circular-reference group(s) in the workbook",
        narrative=(f"Example cycle: {' -> '.join(grp[:5])}. Circular logic with iterative calculation can "
                   "converge to different values depending on calc settings."),
        management_question="Is iterative calculation intentional here, and what anchors the loop?",
        evidence=[CellCitation(ref=r, value=None) for r in grp[:6]],
        confidence=0.8,
    ))


def _rule_working_capital(ctx: FCtx, cands: list[Candidate]) -> None:
    pres = ctx.metrics.scalar("working_capital_presence")
    if pres is None or pres > 0:
        return
    s30 = ctx.metrics.get("receivables_scenario_30d")
    s60 = ctx.metrics.get("receivables_scenario_60d")
    units = (s30.units if s30 else "") or ""
    # cash-accounting declarations in documentation strengthen the evidence
    doc_ev: list[CellCitation] = []
    for name, sd in ctx.wbd.sheets.items():
        for (r, c), rec in sd.cells.items():
            if isinstance(rec.value, str) and re.search(r"cash accounting|cash basis", rec.value, re.I):
                from .ingest import a1
                doc_ev.append(CellCitation(ref=format_ref(name, a1(r, c)), value=None,
                                           label=rec.value.strip()[:120]))
                break
        if doc_ev:
            break
    impact = ""
    if s30 is not None and s30.scalar:
        impact = (f" On the model's own revenue, 30-day receivables would tie up ~{s30.scalar:,.0f} {units}"
                  + (f" and 60-day ~{s60.scalar:,.0f}" if s60 is not None and s60.scalar else "")
                  + " of additional peak cash (tool-computed scenario).")
    cands.append(Candidate(
        category="working_capital_absence",
        severity=Severity.medium,
        title="No working capital anywhere: revenue collects instantly",
        narrative=("The model carries no receivables, payables, or inventory — cash arrives the moment "
                   "revenue books." + impact),
        management_question="What payment terms do customers actually sign, and who finances the gap?",
        evidence=doc_ev,
        confidence=0.85,
    ))


def _rule_valuation_benchmark(ctx: FCtx, cands: list[Candidate]) -> None:
    bm = ctx.benchmarks or {}
    mult = find_scalar_assumptions(ctx.wbd, ctx.structure,
                                   r"ebitda\s*multiple|multiple.*ebitda", (0.5, 100))
    if not mult:
        return
    ref, v, label = mult[0]
    archetype = "hardware_industrial" if (ctx.mapping.get("capacity") or ctx.mapping.get_all("units_sold")) \
        else "default"
    rng = (bm.get("ev_ebitda_multiple") or {}).get(archetype) or (bm.get("ev_ebitda_multiple") or {}).get("default")
    if not rng or v <= rng.get("high", 1e9):
        return
    cands.append(Candidate(
        category="valuation_multiple_above_benchmark",
        severity=Severity.medium,
        title=f"External benchmark (not from workbook): {v:g}x EBITDA exceeds the "
              f"{archetype.replace('_', ' ')} range ({rng['low']:g}–{rng['high']:g}x)",
        narrative=(f"The model's exit multiple ({label!r} = {v:g}x, {ref}) sits above the benchmark range "
                   f"for {archetype.replace('_', ' ')} businesses "
                   f"(External benchmark (not from workbook): {rng['low']:g}–{rng['high']:g}x, typical "
                   f"{rng['typical']:g}x). Software-tier multiples on industrial economics is the single "
                   "highest-leverage assumption in most exit math."),
        management_question=f"Which comparable transactions support {v:g}x for this business model?",
        evidence=[CellCitation(ref=ref, value=v, label=label)],
        confidence=0.7,
    ))


def _rule_convention_violations(ctx: FCtx, cands: list[Candidate]) -> None:
    conv = ctx.structure.input_convention
    if not conv.detected or conv.consistent is not False or not conv.violations:
        return
    cands.append(Candidate(
        category="input_convention_inconsistent",
        severity=Severity.info,
        title=f"{len(conv.violations)}+ typed numbers hide inside formula rows",
        narrative=(f"The workbook marks inputs with {conv.fill_key}"
                   + (f" (declared at {conv.declared_in})" if conv.declared_in else "") +
                   f", but typed constants without that marking sit mid-formula-row, e.g. "
                   f"{', '.join(conv.violations[:5])}. Hard-coded overrides inside computed rows are "
                   "where models quietly diverge from their own logic."),
        management_question="Are these hard-coded cells deliberate overrides, and what do they replace?",
        evidence=[CellCitation(ref=r, value=None) for r in conv.violations[:8]],
        confidence=0.7,
    ))


def _rule_supply_driven(ctx: FCtx, cands: list[Candidate]) -> None:
    ms = ctx.item("market_share")
    if ms is None or ms.n_formula > 0:
        share_inputs = find_scalar_assumptions(ctx.wbd, ctx.structure,
                                               r"penetration|share capture|capture rate|sell.?through", (0.0, 1.0))
        if ms is not None and ms.n_formula > 0 and not share_inputs:
            return
        if ms is None and not share_inputs:
            return
        if share_inputs:
            ref, v, label = share_inputs[0]
            cands.append(Candidate(
                category="supply_driven_revenue",
                severity=Severity.medium,
                title="Market capture is a typed input, not a result",
                narrative=(f"Penetration/sell-through ({label!r} = {v:g}, {ref}) is typed directly rather than "
                           "derived from any demand object — revenue is whatever share is assumed."),
                management_question="What demand evidence (orders, pipeline, design wins) supports the typed capture rate?",
                evidence=[CellCitation(ref=ref, value=v, label=label)],
                confidence=0.65,
            ))
        return
    cands.append(Candidate(
        category="supply_driven_revenue",
        severity=Severity.medium,
        title="Market share is typed directly into the model",
        narrative=(f"The market-share row ({ms.label.strip()!r} on {ms.sheet}) is typed input, not computed "
                   "from demand. Revenue therefore follows from assumed capture, with no demand object "
                   "anywhere to push back."),
        management_question="What converts assumed share into orders — and what does the demand side say?",
        evidence=[CellCitation(ref=ms.refs[-1], value=ms.values[-1] if ms.values else None,
                               label=ms.label.strip())],
        confidence=0.65,
    ))


def _rule_cost_tier_mismatch(ctx: FCtx, cands: list[Candidate]) -> None:
    """A product priced at one tier but costed at a cheaper tier the model's
    own cost sheet distinguishes. Best-effort: needs a cost sheet with >=2
    cost rows for the same product token and a COGS chain into one of them."""
    if ctx.graph is None:
        return
    cogs = ctx.item("cogs_total")
    if cogs is None:
        return
    cost_sheets = [name for name, (role, _) in ctx.structure.roles.items()
                   if role == SheetRole.cost_buildup]
    if not cost_sheets:
        return
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
    cogs_up: set[int] = set()
    for ref in cogs.refs[:4] + cogs.refs[-2:]:
        s, cell = parse_ref(ref)
        cl, row = coordinate_from_string(cell)
        cogs_up |= ctx.graph.upstream(ctx.graph.node_of(s, row, column_index_from_string(cl)), cap=60_000)

    for cs in cost_sheets:
        sd = ctx.wbd.sheets[cs]
        si = ctx.graph.sheet_index[cs]
        # gather labeled numeric rows: token -> [(row, label, value, in_cogs_chain)]
        rows: dict[int, str] = {}
        for (r, c), rec in sd.cells.items():
            if isinstance(rec.value, str) and 2 < len(rec.value.strip()) < 80 and c <= 4:
                rows.setdefault(r, rec.value.strip())
        by_token: dict[str, list[tuple[int, str, float, bool]]] = {}
        for r, lab in rows.items():
            if not re.search(r"cost|price|\$", lab, re.I):
                continue
            vals = [(c, rec) for (rr, c), rec in sd.cells.items()
                    if rr == r and is_number(rec.value)]
            if not vals:
                continue
            c0, rec0 = sorted(vals)[0]
            in_chain = any((node >> 44) == si and ((node >> 20) & ((1 << 24) - 1)) == r
                           for node in cogs_up if node >= 0)
            token = re.sub(r"\b(cost|price|per|unit|of|the|a|an|\$|/)\b", " ", lab.lower())
            token = re.sub(r"[^a-z]+", " ", token).strip().split(" ")[0] if token.strip() else ""
            if token:
                by_token.setdefault(token, []).append((r, lab, float(rec0.value), in_chain))
        for token, entries in by_token.items():
            if len(entries) < 2:
                continue
            used = [e for e in entries if e[3]]
            unused = [e for e in entries if not e[3]]
            if not used or not unused:
                continue
            cheap = min(entries, key=lambda e: e[2])
            dear = max(entries, key=lambda e: e[2])
            if dear[2] <= cheap[2] * 1.15 or not cheap[3] or dear[3]:
                continue
            cands.append(Candidate(
                category="cost_tier_mismatch",
                severity=Severity.critical,
                title=f"COGS pulls the cheaper {token!r} cost tier while a pricier tier exists",
                narrative=(f"The cost sheet {cs!r} defines at least two cost levels for {token!r}: "
                           f"{cheap[1]!r} = {cheap[2]:,.4g} (feeds COGS) and {dear[1]!r} = {dear[2]:,.4g} "
                           f"(does not). If the product is priced/spec'd at the {dear[1]!r} level but costed "
                           f"at {cheap[1]!r}, gross margin is overstated by construction."),
                management_question=(f"Which tier is actually being sold — and if it is {dear[1]!r}, "
                                     "why does COGS reference the cheaper row?"),
                evidence=[
                    CellCitation(ref=format_ref(cs, f"row {cheap[0]}"), value=cheap[2], label=cheap[1]),
                    CellCitation(ref=format_ref(cs, f"row {dear[0]}"), value=dear[2], label=dear[1]),
                ],
                confidence=0.6,
                acquit_with=("tier_in_price_chain",),
            ))
            return  # one per workbook; this rule is noisy by nature


RULES = [
    _rule_tie_out_failures, _rule_missing_cached, _rule_units_inconsistency,
    _rule_grant_dependence, _rule_growth_whiplash, _rule_hockey_stick,
    _rule_financing_cliff, _rule_concentration, _rule_market_share_basis,
    _rule_frozen_costs, _rule_expense_decoupling, _rule_capex_in_opex,
    _rule_opex_pegged_to_revenue, _rule_flat_inputs, _rule_one_way_margin,
    _rule_taxes, _rule_orphans_and_errors, _rule_circulars,
    _rule_working_capital, _rule_valuation_benchmark,
    _rule_convention_violations, _rule_supply_driven, _rule_cost_tier_mismatch,
]


# ===========================================================================
# Phase 6 — acquittal pass
# ===========================================================================

def _try_acquit(ctx: FCtx, cand: Candidate) -> Optional[Acquittal]:
    P = ctx.periods()

    if "small_base" in cand.acquit_with and cand.periods:
        it = ctx.item("revenue_total")
        if it is not None:
            levels = {p: it.value_for(p) for p in P}
            scale = max((abs(v) for v in levels.values() if v is not None), default=0.0)
            bases = []
            for p in cand.periods:
                i = P.index(p)
                prev = levels.get(P[i - 1]) if i > 0 else None
                if prev is not None:
                    bases.append((P[i - 1], prev))
            if bases and all(abs(b) < SMALL_BASE_FRACTION * scale for _, b in bases):
                # growth ON a tiny base is noise — unless the candidate is about
                # a decline being smoothed (handled by the grant rule, not here)
                return Acquittal(
                    candidate=cand.title,
                    category=cand.category,
                    decomposition=(f"The growth bases are tiny relative to terminal scale "
                                   f"({'; '.join(f'{p}: {b:,.0f}' for p, b in bases)} vs series max "
                                   f"{scale:,.0f}). Pre-product noise, not a trajectory claim."),
                    resolution_cells=[it.ref_for(p) for p, _ in bases if it.ref_for(p)],
                )

    if "grant_timing" in cand.acquit_with and cand.periods:
        share = ctx.metrics.series("grant_share_of_revenue")
        exg = ctx.metrics.series("revenue_ex_grants_growth")
        if share and exg:
            p = cand.periods[-1]
            sh_now = share.get(p)
            i = P.index(p)
            sh_prev = share.get(P[i - 1]) if i > 0 else None
            ex = exg.get(p)
            if (sh_now is not None and sh_prev is not None and abs(sh_now - sh_prev) > 0.2
                    and ex is not None and abs(ex) < 1.0):
                m = ctx.metrics.get("grant_share_of_revenue")
                return Acquittal(
                    candidate=cand.title,
                    category=cand.category,
                    decomposition=(f"The swing is grant timing: grant share moves {sh_prev:.0%} -> {sh_now:.0%} "
                                   f"into {p} while ex-grant growth is an unremarkable {ex:+.0%}. The anomaly "
                                   "belongs to the grant rows, not the growth row."),
                    resolution_cells=m.inputs[:6] if m else [],
                )

    if "ex_capex" in cand.acquit_with and cand.periods:
        exg = ctx.metrics.series("opex_ex_capex_growth")
        if exg:
            p = cand.periods[0]
            v = exg.get(p)
            if v is not None and v > 0.0:
                m = ctx.metrics.get("opex_ex_capex_growth")
                return Acquittal(
                    candidate=cand.title,
                    category=cand.category,
                    decomposition=(f"Ex-capital-equipment, expenses GROW {v:+.0%} in {p} — the decline is "
                                   "capex lumpiness (equipment purchases rolling off), not costs being "
                                   "deleted while the business multiplies."),
                    resolution_cells=m.inputs[:8] if m else [],
                )

    if "tax_trend" in cand.acquit_with:
        eff = ctx.metrics.series("effective_tax_rate")
        stat = ctx.metrics.scalar("statutory_tax_rate_typed")
        if eff and stat:
            vals = [(p, v) for p, v in eff.items() if v is not None]
            if len(vals) >= 3 and vals[0][1] < stat * 0.8 and abs(vals[-1][1] - stat) < stat * 0.25:
                return Acquittal(
                    candidate=cand.title,
                    category=cand.category,
                    decomposition=(f"The effective rate climbs from {vals[0][1]:.0%} toward the typed statutory "
                                   f"{stat:.0%} — the path is consistent with carryforwards burning off."),
                    resolution_cells=[],
                )

    return None


# ===========================================================================
# Phase 7 — severity, dedup, cause-tracing, assembly
# ===========================================================================

def _node_of_ref(ctx: FCtx, ref: str) -> Optional[int]:
    try:
        from openpyxl.utils.cell import column_index_from_string, coordinate_from_string
        s, cell = parse_ref(ref)
        cl, row = coordinate_from_string(cell)
        return ctx.graph.node_of(s, row, column_index_from_string(cl))
    except Exception:
        return None


_PRIORITY = ["arithmetic_integrity", "data_quality", "unfalsifiable_check", "cost_tier_mismatch",
             "unit_scale_inconsistency", "grant_dependence", "capex_expensed_in_opex_orphaned_depreciation",
             "market_share_basis_mismatch", "financing_cliff", "frozen_support_costs",
             "expense_growth_decoupled", "concentration", "capex_expensed_in_opex"]


def _dedup_echoes(ctx: FCtx, cands: list[Candidate]) -> list[Candidate]:
    """Collapse downstream symptoms onto the finding closest to the source."""
    if ctx.graph is None or len(cands) < 2:
        return cands
    keep: list[Candidate] = []
    absorbed: set[int] = set()
    order = sorted(range(len(cands)),
                   key=lambda i: _PRIORITY.index(cands[i].category) if cands[i].category in _PRIORITY else 99)
    for i in order:
        if i in absorbed:
            continue
        cause = cands[i]
        cause_nodes = {n for n in (_node_of_ref(ctx, r) for r in cause.cells()) if n is not None}
        if not cause_nodes:
            keep.append(cause)
            continue
        downstream: set[int] = set()
        for n in list(cause_nodes)[:6]:
            downstream |= ctx.graph.downstream(n, cap=30_000)
        for j in range(len(cands)):
            if j == i or j in absorbed or cands[j].category == cause.category:
                continue
            echo = cands[j]
            if echo.category in _PRIORITY:
                continue  # primary-story categories are never absorbed as echoes
            echo_nodes = {n for n in (_node_of_ref(ctx, r) for r in echo.cells()) if n is not None}
            if echo_nodes and echo_nodes <= downstream:
                cause.narrative += f" (Downstream symptom folded in: {echo.title}.)"
                cause.evidence.extend(echo.evidence[:2])
                absorbed.add(j)
        keep.append(cause)
    for j in range(len(cands)):
        if j not in absorbed and cands[j] not in keep:
            keep.append(cands[j])
    return keep


def build_findings(ctx: FCtx) -> FindingsResult:
    res = FindingsResult()
    cands: list[Candidate] = []
    for rule in RULES:
        try:
            rule(ctx, cands)
        except Exception:
            log.exception("finding rule %s failed", rule.__name__)
            res.notes.append(f"finding rule {rule.__name__} failed and was skipped")

    # Phase 6: acquittal pass
    survivors: list[Candidate] = []
    for cand in cands:
        acq = None
        try:
            acq = _try_acquit(ctx, cand)
        except Exception:
            log.exception("acquittal attempt failed for %s", cand.category)
        if acq is not None:
            res.acquittals.append(acq)
        else:
            survivors.append(cand)

    # Phase 7: echo dedup, severity caps, assembly
    survivors = _dedup_echoes(ctx, survivors)
    trust_degraded = ctx.integrity.trust_degraded
    sev_rank = {Severity.critical: 0, Severity.high: 1, Severity.medium: 2, Severity.info: 3}
    survivors.sort(key=lambda c: (sev_rank[c.severity], -c.confidence))
    for n, cand in enumerate(survivors, 1):
        sev = cand.severity
        conf = cand.confidence
        if cand.mapping_confidence < 0.6 and sev in (Severity.critical, Severity.high):
            sev = Severity.medium
            cand.narrative += (f" (Severity capped: the underlying row mapping carries only "
                               f"{cand.mapping_confidence:.2f} confidence.)")
            conf = min(conf, cand.mapping_confidence)
        res.findings.append(Finding(
            id=f"F{n}",
            category=cand.category,
            severity=sev,
            title=cand.title,
            narrative=cand.narrative,
            evidence_cells=[e.ref for e in cand.evidence],
            evidence_values=cand.evidence,
            quantified_impact=cand.quantified_impact,
            management_question=cand.management_question,
            confidence=round(conf, 2),
            trust_degraded=trust_degraded and cand.category != "arithmetic_integrity",
        ))

    # clean checks from passing tie-outs
    for t in ctx.integrity.tie_outs:
        if t.status == TieOutStatus.passed:
            res.clean_checks.append(CleanCheck(
                check=t.label or t.id,
                result=t.detail or f"passes across {t.periods_checked} periods",
                evidence=[e.ref for e in t.evidence][:6],
            ))
    if ctx.graph is not None and not ctx.graph.ref_error_formula_cells and not ctx.graph.error_nodes:
        res.clean_checks.append(CleanCheck(
            check="No error cells anywhere in the workbook",
            result="no #REF!/#DIV/0!/#VALUE! values or formulas found",
            evidence=[],
        ))
    conv = ctx.structure.input_convention
    if conv.detected and conv.consistent:
        res.clean_checks.append(CleanCheck(
            check="Input color convention applied consistently",
            result=f"typed inputs uniformly carry {conv.fill_key}; no hard-codes hiding in formula rows",
            evidence=[conv.declared_in] if conv.declared_in else [],
        ))
    return res
